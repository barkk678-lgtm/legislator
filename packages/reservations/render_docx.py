"""פלט Word של ההסתייגויות (הסתייגויות 5, 26.9).

ההנחיות של ברק:
- **David 12, מיושר לשני הצדדים.** בלי עיצוב מעבר לזה.
- כותרת: שם ההצעה.
- **כותרות מיקום, כמו בטבריה:** "לפני סעיף 1", "לסעיף 1", "לאחרי סעיף 1" -
  בלעדיהן "בפסקה (2)" לא אומר כלום.
- **כל הסתייגות במספר, ברצף אחד** על פני כל הכותרות - 1, 2, 3... **בלי חלופות
  באותיות** (הקיבוץ לחלופות הוא עבודה של היועצים המשפטיים).
- **בלי השורה "קבוצת X מציעה:".**
- עובר את בודק המבנה (packages/render/docx_check) - בלי אזהרה בפתיחה.

המספור נכתב כטקסט ("1. ") ולא כרשימה אוטומטית של Word: מספור אוטומטי מתאפס
או משתבש כשמעתיקים קטעים, והמספר הוא חלק מההסתייגות (ההצבעה היא לפי מספר).
"""

from __future__ import annotations

import io
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

FONT = "David"
SIZE = Pt(12)


_CONTROL = re.compile(r"[\x00-\x08\x0c\x0e-\x1f]")


def _xml_safe(text: str) -> str:
    """ירידת שורה ידנית של Word (\x0b) - רווח; שאר תווי הבקרה - נמחקים. בלעדיהם
    python-docx קורס (ValueError), ו-XML 1.0 אוסר אותם (tests/unit/test_docx_structure.py)."""
    return _CONTROL.sub("", text.replace("\x0b", " ").replace("\x07", ""))


def _rtl_paragraph(doc, text: str):
    text = _xml_safe(text)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    ppr = p._p.get_or_add_pPr()
    ppr.insert(0, OxmlElement("w:bidi"))
    run = p.add_run(text)
    run._r.get_or_add_rPr().append(OxmlElement("w:rtl"))
    return p


def _set_default_font(doc) -> None:
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = SIZE
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), FONT)
    szcs = OxmlElement("w:szCs")
    szcs.set(qn("w:val"), str(int(SIZE.pt * 2)))
    rpr.append(szcs)
    lang = OxmlElement("w:lang")
    lang.set(qn("w:bidi"), "he-IL")
    rpr.append(lang)


def build_document(*, bill_title: str, items: list) -> Document:
    """items: families.Item (או כל אובייקט עם heading ו-lines), בסדר המסמך."""
    doc = Document()
    _set_default_font(doc)
    _rtl_paragraph(doc, bill_title or "הצעת חוק")
    current = None
    for number, item in enumerate(items, 1):
        if item.heading != current:
            current = item.heading
            _rtl_paragraph(doc, current)
        first, *rest = item.lines
        _rtl_paragraph(doc, f"{number}. {first}")
        for line in rest:
            _rtl_paragraph(doc, line)
    if not items:
        _rtl_paragraph(doc, "אין הסתייגויות")
    return doc


def docx_bytes(*, bill_title: str, items: list) -> bytes:
    buf = io.BytesIO()
    build_document(bill_title=bill_title, items=items).save(buf)
    return buf.getvalue()


__all__ = ["build_document", "docx_bytes"]
