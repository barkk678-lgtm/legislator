"""פלט ה-Word של ההסתייגויות (הסתייגויות 5, 26.9).

David 12, מיושר לשני הצדדים; כותרת - שם ההצעה; כותרות מיקום כמו בטבריה ("לסעיף 1",
"לאחרי סעיף 3") פעם אחת לכל קבוצה; מספור רציף 1, 2, 3 על פני כל הכותרות, בלי
חלופות באותיות; בלי "קבוצת X מציעה:"; ועובר את בודק המבנה (docx_check).
"""

import io
import re
import sys
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "packages" / "reservations"))
sys.path.insert(0, str(ROOT / "packages" / "render"))

from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx_check import structural_problems  # noqa: E402
import families  # noqa: E402
import render_docx  # noqa: E402
from pdf_bill import parse_bill_pdf  # noqa: E402

BILL = parse_bill_pdf(ROOT / "reference" / "הצעת חוק טבריה.pdf")
ITEMS = families.plan(BILL, "serious", list(families.FAMILIES), 1000).items
DATA = render_docx.docx_bytes(bill_title=BILL.title, items=ITEMS)
DOC = Document(io.BytesIO(DATA))
PARAS = [p.text for p in DOC.paragraphs]


def test_title_is_the_bill():
    assert PARAS[0] == BILL.title, PARAS[0]


def test_location_headings_once_per_group_and_in_order():
    headings = [p for p in PARAS if re.fullmatch(r"(?:לפני |לאחרי |ל)סעיף \d+[א-ת]?", p)]
    assert headings == list(dict.fromkeys(it.heading for it in ITEMS)), headings
    assert headings[0] == "לסעיף 1", headings


def test_continuous_numbering_no_letters():
    numbered = [p for p in PARAS if re.match(r"^\d+\. ", p)]
    assert [int(p.split(".")[0]) for p in numbered] == list(range(1, len(ITEMS) + 1))
    assert not any(re.match(r"^[א-ת]\. ", p) for p in PARAS), "חלופות באותיות"


def test_no_group_line():
    assert not any("מציעה" in p or "מציעים" in p for p in PARAS)


def test_david_12_justified_rtl():
    normal = DOC.styles["Normal"]
    assert normal.font.name == "David" and normal.font.size.pt == 12
    body = [p for p in DOC.paragraphs if p.text]
    assert all(p.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY for p in body)
    xml = zipfile.ZipFile(io.BytesIO(DATA)).read("word/document.xml").decode()
    assert xml.count("<w:bidi/>") >= len(body) and "<w:rtl/>" in xml


def test_passes_structure_check():
    assert structural_problems(DATA) == [], structural_problems(DATA)


def test_empty_and_control_chars():
    empty = Document(io.BytesIO(render_docx.docx_bytes(bill_title="", items=[])))
    assert [p.text for p in empty.paragraphs] == ["הצעת חוק", "אין הסתייגויות"]
    item = families.Item(heading="לסעיף 1", lines=["במקום \"א\x0bב\" יבוא \"ג\x07\"."], family="x",
                         group="g", order=(0,))
    data = render_docx.docx_bytes(bill_title="חוק", items=[item])
    assert structural_problems(data) == []
    assert Document(io.BytesIO(data)).paragraphs[2].text == '1. במקום "א ב" יבוא "ג".'


def test_a4_page():
    """הכרעה ח (ברק, 26.9.2026): A4 - 11906×16838 twips, לא Letter (12240×15840)."""
    xml = zipfile.ZipFile(io.BytesIO(DATA)).read("word/document.xml").decode("utf-8")
    sizes = re.findall(r'<w:pgSz [^>]*w:w="(\d+)"[^>]*w:h="(\d+)"', xml)
    assert sizes == [("11906", "16838")], sizes


for name, fn in sorted(list(globals().items())):
    if name.startswith("test_"):
        fn()
        print(f"  ✓ {name}")
print("test_reservations_render: כל הבדיקות עברו")
